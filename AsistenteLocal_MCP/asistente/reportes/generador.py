"""Generador de reportes/documentos **locales** (Markdown y Word .docx).

Los reportes se escriben en `data/reportes/`. No se envían a Dataverse.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from ..config import config

try:
    from docx import Document
except Exception:  # pragma: no cover - dependencia opcional
    Document = None


def _timestamp() -> str:
    return dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def _ruta(nombre: str, ext: str) -> Path:
    config.reportes_dir.mkdir(parents=True, exist_ok=True)
    return config.reportes_dir / f"{nombre}.{ext}"


def generar_markdown(
    titulo: str,
    secciones: list[dict],
    nombre: str | None = None,
) -> Path:
    """Genera un reporte en Markdown.

    `secciones`: lista de {"cabecera": str, "cuerpo": str (markdown)}.
    """
    nombre = nombre or f"reporte_{_timestamp()}"
    lineas = [f"# {titulo}", "", f"_Generado: {dt.datetime.now().isoformat(timespec='minutes')}_", ""]
    for sec in secciones:
        lineas.append(f"## {sec.get('cabecera', '')}")
        lineas.append("")
        lineas.append(sec.get("cuerpo", ""))
        lineas.append("")
    path = _ruta(nombre, "md")
    path.write_text("\n".join(lineas), encoding="utf-8")
    return path


def generar_docx(
    titulo: str,
    secciones: list[dict],
    nombre: str | None = None,
) -> Path:
    """Genera un reporte en Word (.docx). Requiere python-docx."""
    if Document is None:
        raise RuntimeError(
            "python-docx no está instalado. Ejecuta: pip install python-docx"
        )
    nombre = nombre or f"reporte_{_timestamp()}"
    doc = Document()
    doc.add_heading(titulo, level=0)
    doc.add_paragraph(f"Generado: {dt.datetime.now().isoformat(timespec='minutes')}")
    for sec in secciones:
        doc.add_heading(sec.get("cabecera", ""), level=1)
        for parrafo in (sec.get("cuerpo", "") or "").split("\n"):
            if parrafo.strip():
                doc.add_paragraph(parrafo.strip())
    path = _ruta(nombre, "docx")
    doc.save(str(path))
    return path


def generar(
    titulo: str,
    secciones: list[dict],
    nombre: str | None = None,
    formatos: list[str] | None = None,
) -> list[Path]:
    """Genera el reporte en los formatos indicados (md/docx)."""
    formatos = formatos or ["markdown"]
    rutas: list[Path] = []
    for fmt in formatos:
        if fmt in ("markdown", "md"):
            rutas.append(generar_markdown(titulo, secciones, nombre))
        elif fmt == "docx":
            rutas.append(generar_docx(titulo, secciones, nombre))
        else:
            raise ValueError(f"Formato no soportado: {fmt}")
    return rutas
